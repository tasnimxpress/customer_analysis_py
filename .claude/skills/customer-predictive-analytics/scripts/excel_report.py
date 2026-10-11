"""Write a consistently formatted Excel results workbook.

Usage:
    from excel_report import write_report
    write_report(
        "analysis/churn_logistic_results.xlsx",
        sheets={"Coefficients": coef_df, "Predictions": pred_df},
        summary={"Accuracy": 0.81, "Baseline accuracy": 0.58},   # optional key metrics
        notes=["Accuracy beats the always-'no churn' baseline by 23 points."],  # optional
        images={"Charts": ["analysis/roc.png", "analysis/cm.png"]},  # optional
    )

The workbook always starts with a "Summary" sheet (metrics + notes), followed by
each DataFrame in `sheets` (header styled, frozen, autofiltered, columns sized),
then any image sheets.
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_table(ws, max_width=50):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), max_width)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000" if abs(cell.value) < 1 else "#,##0.00"


def write_report(path, sheets, summary=None, notes=None, images=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    summary_df = pd.DataFrame(
        list((summary or {}).items()), columns=["Metric", "Value"]
    )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)

    wb = load_workbook(path)
    for name in ["Summary", *[n[:31] for n in sheets]]:
        _style_table(wb[name])

    if notes:
        ws = wb["Summary"]
        start = ws.max_row + 2
        ws.cell(row=start, column=1, value="Notes").font = Font(bold=True, size=12)
        for i, note in enumerate(notes, start=1):
            c = ws.cell(row=start + i, column=1, value=f"• {note}")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=start + i, start_column=1, end_row=start + i, end_column=6)
            ws.row_dimensions[start + i].height = 30

    for sheet_name, paths in (images or {}).items():
        ws = wb.create_sheet(sheet_name[:31])
        row = 1
        for img_path in paths:
            img = XLImage(str(img_path))
            ws.add_image(img, f"A{row}")
            row += int(img.height / 20) + 2

    wb.save(path)
    return path
